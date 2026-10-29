"""DOCX. Paragraphs and tables in document order.

Tables matter: a payment schedule rendered as a table is exactly the thing an
attorney asks about, and dropping it would make the answer "not in your
documents" when it demonstrably is.
"""

from __future__ import annotations

from app.ingestion.ports import ParsedPage, ParseResult

CHARS_PER_PAGE = 2600


def _table_to_text(table: object) -> str:
    rows: list[str] = []
    for row in getattr(table, "rows", []):
        cells = [" ".join(c.text.split()) for c in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


class DocxParser:
    name = "docx/python-docx"

    def supports(self, ext: str) -> bool:
        return ext == ".docx"

    def parse(self, path: str, *, max_pages: int) -> ParseResult:
        import docx
        from docx.document import Document as DocxDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document: DocxDocument = docx.Document(path)
        blocks: list[str] = []

        # Walk the body in order so a table stays next to the clause introducing
        # it. `document.paragraphs` and `document.tables` separately would lose
        # that, and with it the reading order a citation depends on.
        body = document.element.body
        for child in body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                text = Paragraph(child, document).text.strip()
                if text:
                    blocks.append(text)
            elif tag == "tbl":
                text = _table_to_text(Table(child, document))
                if text:
                    blocks.append(text)

        pages: list[ParsedPage] = []
        buf: list[str] = []
        size = 0
        for block in blocks:
            if size + len(block) > CHARS_PER_PAGE and buf:
                pages.append(
                    ParsedPage(page_no=len(pages) + 1, text="\n\n".join(buf), engine=self.name)
                )
                buf, size = [], 0
                if len(pages) >= max_pages:
                    break
            buf.append(block)
            size += len(block) + 2
        if buf and len(pages) < max_pages:
            pages.append(
                ParsedPage(page_no=len(pages) + 1, text="\n\n".join(buf), engine=self.name)
            )
        if not pages:
            pages = [ParsedPage(page_no=1, text="", engine=self.name)]

        return ParseResult(pages=pages, scanned=False, engine=self.name)
